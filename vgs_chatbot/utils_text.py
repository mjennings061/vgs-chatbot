"""Document parsing and chunking utilities."""

from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass
from typing import List, Sequence, Tuple

import pdfplumber
from docx import Document

logger = logging.getLogger(__name__)


@dataclass
class TextBlock:
    """Atomic block of text prior to chunking."""

    kind: str  # "text" or "table"
    text: str
    page_start: int | None = None
    page_end: int | None = None


@dataclass
class Section:
    """Structured section detected on a page."""

    title: str
    blocks: List[TextBlock]
    order_code: str | None = None


@dataclass
class ChunkFragment:
    """Chunk emitted after block-aware segmentation."""

    text: str
    kind: str  # "text" or "table"
    page_start: int | None = None
    page_end: int | None = None


def read_pdf(file_bytes: bytes) -> List[Tuple[int, str]]:
    """Extract plain text from each page of a PDF.

    Args:
        file_bytes: Raw bytes representing the PDF file.

    Returns:
        list[tuple[int, str]]: Sequence of page numbers and their extracted text.
    """
    pages: List[Tuple[int, str]] = []
    text_lengths: List[int] = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for index, page in enumerate(pdf.pages, start=1):
            text = page.extract_text() or ""
            stripped = text.strip()
            pages.append((index, stripped))
            text_lengths.append(len(stripped))
    logger.debug("Extracted text from %s PDF pages.", len(pages))
    _assert_pdf_has_text(pages, text_lengths)
    return pages


def read_docx(file_bytes: bytes) -> List[Tuple[int, str]]:
    """Extract text from a Word document, approximating pages.

    Args:
        file_bytes: Raw bytes representing the Word document.

    Returns:
        list[tuple[int, str]]: Approximate page number and associated text.
    """
    doc = Document(io.BytesIO(file_bytes))
    paragraphs = [para.text.strip() for para in doc.paragraphs if para.text.strip()]
    if not paragraphs:
        return [(1, "")]
    # Word does not expose pagination without layout; treat blocks of ~30 paragraphs as a page.
    chunk_size = 30
    pages: List[Tuple[int, str]] = []
    for index in range(0, len(paragraphs), chunk_size):
        page_no = len(pages) + 1
        page_text = "\n".join(paragraphs[index : index + chunk_size])
        pages.append((page_no, page_text))
    logger.debug(
        "Extracted %s DOCX paragraphs across %s pseudo-pages.",
        len(paragraphs),
        len(pages),
    )
    return pages


def detect_sections(text: str) -> List[Section]:
    """Split document text into structured sections with table-aware blocks.

    This function normalises noisy line breaks, removes headers/footers, detects
    headings (including order codes like DHO/AESO/GASO), and preserves table-like
    blocks so they are not fragmented mid-row. Page markers of the form
    `[[[PAGE_BREAK_<n>]]]` are honoured so page ranges propagate to chunks.

    Args:
        text: Document-level text (may include multiple pages) to analyse.

    Returns:
        list[Section]: Detected section titles with block content.
    """
    lines = _prepare_lines(text)
    if not lines:
        return []

    sections: List[Section] = []
    current_title: str | None = None
    current_blocks: List[TextBlock] = []
    paragraph_buffer: List[tuple[str, int | None]] = []
    table_buffer: List[tuple[str, int | None]] = []
    current_page: int | None = None
    seen_pages: set[int] = set()

    def flush_paragraph() -> None:
        if paragraph_buffer:
            merged = _join_paragraph_lines([p for p, _ in paragraph_buffer])
            if merged:
                pages = [p for _, p in paragraph_buffer if p is not None]
                page_start = min(pages) if pages else current_page
                page_end = max(pages) if pages else current_page
                current_blocks.append(
                    TextBlock(
                        kind="text",
                        text=merged,
                        page_start=page_start,
                        page_end=page_end,
                    )
                )
            paragraph_buffer.clear()

    def flush_table() -> None:
        if table_buffer:
            table_text = "\n".join([p for p, _ in table_buffer]).strip()
            if table_text:
                pages = [p for _, p in table_buffer if p is not None]
                page_start = min(pages) if pages else current_page
                page_end = max(pages) if pages else current_page
                current_blocks.append(
                    TextBlock(
                        kind="table",
                        text=table_text,
                        page_start=page_start,
                        page_end=page_end,
                    )
                )
            table_buffer.clear()

    def flush_section() -> None:
        nonlocal current_blocks, current_title
        flush_table()
        flush_paragraph()
        if current_blocks:
            sections.append(
                Section(
                    title=clean_title(current_title or "General") or "General",
                    blocks=current_blocks,
                    order_code=_extract_order_code(current_title),
                )
            )
            current_blocks = []

    for index, line in enumerate(lines):
        next_line = lines[index + 1] if index + 1 < len(lines) else None
        page_marker = _extract_page_marker(line)
        if page_marker is not None:
            current_page = page_marker
            seen_pages.add(page_marker)
            continue

        if not line.strip():
            flush_table()
            flush_paragraph()
            continue

        if _looks_like_heading(line, next_line):
            flush_section()
            current_title = clean_title(line)
            continue

        if _looks_like_table_line(line):
            flush_paragraph()
            table_buffer.append((line, current_page))
            continue

        if _is_bullet_line(line):
            flush_paragraph()
            paragraph_buffer.append((line, current_page))
            flush_paragraph()
            continue

        # Join wrapped lines unless this is the first line of a new paragraph.
        if paragraph_buffer:
            prev, prev_page = paragraph_buffer[-1]
            if prev.endswith("-"):
                paragraph_buffer[-1] = (f"{prev[:-1]}{line}", prev_page or current_page)
            else:
                paragraph_buffer[-1] = (f"{prev} {line}", prev_page or current_page)
        else:
            paragraph_buffer.append((line, current_page))

    flush_section()
    if not sections and text.strip():
        page_start = min(seen_pages) if seen_pages else None
        page_end = max(seen_pages) if seen_pages else None
        sections = [
            Section(
                title="General",
                blocks=[
                    TextBlock(
                        kind="text",
                        text=_join_paragraph_lines(lines),
                        page_start=page_start,
                        page_end=page_end,
                    )
                ],
                order_code=_extract_order_code(text),
            )
        ]
    logger.debug("Detected %s sections.", len(sections))
    return sections


def chunk_text(
    text: str | None = None,
    *,
    blocks: Sequence[TextBlock] | None = None,
    target_chars: int = 900,
    overlap: int = 120,
) -> List[ChunkFragment]:
    """Chunk text into roughly target-sized pieces with table protection.

    Long text blocks are split sentence-first (when possible) to avoid mid-sentence
    breaks, then fall back to overlapping character windows. Table blocks are kept
    intact and only split between rows if they exceed the target size.

    Args:
        text: Text to segment (optional when passing blocks directly).
        blocks: Pre-tokenised blocks preserving table boundaries.
        target_chars: Preferred chunk size in characters.
        overlap: Overlap length applied when splitting long paragraphs.

    Returns:
        list[ChunkFragment]: Ordered chunk fragments with kind metadata.
    """
    if blocks is None:
        base_text = (text or "").strip()
        if not base_text:
            return []
        blocks = [TextBlock(kind="text", text=base_text)]

    chunks: List[ChunkFragment] = []
    buffer = ""
    buffer_pages: set[int] = set()

    for block in blocks:
        if block.kind == "table":
            if buffer:
                chunks.append(
                    ChunkFragment(
                        text=buffer.strip(),
                        kind="text",
                        page_start=min(buffer_pages) if buffer_pages else block.page_start,
                        page_end=max(buffer_pages) if buffer_pages else block.page_end,
                    )
                )
                buffer = ""
                buffer_pages.clear()
            for table_chunk in _split_table_block(block.text, target_chars):
                if table_chunk:
                    chunks.append(
                        ChunkFragment(
                            text=table_chunk,
                            kind="table",
                            page_start=block.page_start,
                            page_end=block.page_end,
                        )
                    )
            continue

        for paragraph in _split_text_block(block.text, target_chars, overlap):
            if not paragraph:
                continue
            if len(buffer) + len(paragraph) + 2 <= target_chars:
                buffer = f"{buffer}\n\n{paragraph}".strip()
                buffer_pages.update(_page_numbers(block))
            else:
                if buffer:
                    chunks.append(
                        ChunkFragment(
                            text=buffer.strip(),
                            kind="text",
                            page_start=min(buffer_pages) if buffer_pages else block.page_start,
                            page_end=max(buffer_pages) if buffer_pages else block.page_end,
                        )
                    )
                buffer = paragraph
                buffer_pages = set(_page_numbers(block))

    if buffer:
        chunks.append(
            ChunkFragment(
                text=buffer.strip(),
                kind="text",
                page_start=min(buffer_pages) if buffer_pages else None,
                page_end=max(buffer_pages) if buffer_pages else None,
            )
        )

    logger.debug("Chunked text into %s segments.", len(chunks))
    return chunks


def _prepare_lines(text: str) -> List[str]:
    """Normalise raw text for section detection."""
    normalised = text.replace("\r\n", "\n").replace("\r", "\n")
    # Heal hyphenated line breaks early to avoid split words.
    normalised = re.sub(r"([A-Za-z])-\s*\n\s*([A-Za-z])", r"\1\2", normalised)
    normalised = re.sub(r"\n{3,}", "\n\n", normalised)
    lines = [line.strip() for line in normalised.splitlines()]
    cleaned: List[str] = []
    for line in lines:
        if _is_header_footer(line):
            continue
        cleaned.append(line)
    return cleaned


def _looks_like_heading(line: str, next_line: str | None = None) -> bool:
    """Return True if a line resembles a heading."""
    if not line:
        return False
    if _looks_like_table_line(line):
        return False

    cleaned = line.strip(" \t-–:·")
    if len(cleaned) < 4 or len(cleaned) > 120:
        return False

    # Very short alphabetic strings like "A O O" are often line-wrapped words, not headings.
    alpha_chars = [ch for ch in cleaned if ch.isalpha()]
    if len(alpha_chars) < 4 and not re.match(
        r"^(?:RA|AESO|GASO|DHO|DHE|DHI)\s*\d{3,4}(?:\([^)]+\))?:?.*$", cleaned, re.IGNORECASE
    ):
        return False

    if re.match(r"^(?:RA|AESO|GASO|DHO|DHE|DHI)\s*\d{3,4}(?:\([^)]+\))?:?.*$", cleaned, re.IGNORECASE):
        return True
    if re.match(r"^\d+(\.\d+)*\s+[A-Z].*", cleaned):
        return True

    words = cleaned.split()
    uppercase = sum(1 for ch in alpha_chars if ch.isupper()) if alpha_chars else 0
    ratio = (uppercase / len(alpha_chars)) if alpha_chars else 0.0
    if ratio > 0.82 and len(words) <= 12:
        return True
    if cleaned.endswith(":") and ratio > 0.6 and len(words) <= 15:
        return True

    # Headings are often surrounded by blank lines; prefer those.
    if ratio > 0.65 and len(words) <= 8 and (next_line is None or not next_line.strip()):
        return True
    return False


def _looks_like_table_line(line: str) -> bool:
    """Heuristic to keep table rows intact."""
    if not line:
        return False
    if "|" in line or "\t" in line:
        return True
    if re.search(r"\s{2,}", line):
        return True

    tokens = line.split()
    if not tokens:
        return False

    x_count = sum(1 for tok in tokens if tok.upper() == "X")
    digit_count = sum(1 for tok in tokens if re.fullmatch(r"\d+(\.\d+)?", tok))
    if x_count >= 2 or (x_count >= 1 and digit_count >= 1):
        return True
    if digit_count >= 2 and len(tokens) <= 12:
        return True

    if re.match(r"^\d{3,4}\([^)]+\):", line):
        return True
    if re.match(r"^(?:AESO|GASO|DHO|DHE|DHI)\s*\d{3,4}\b", line, re.IGNORECASE):
        return True
    return False


def _is_header_footer(line: str) -> bool:
    """Return True for boilerplate headers/footers that fragment chunks."""
    if not line:
        return False
    if re.search(r"UNCONTROLLED COPY WHEN PRINTED", line, re.IGNORECASE):
        return True
    if re.match(r"Page\s+\d+\s+of\s+\d+", line, re.IGNORECASE):
        return True
    if re.search(r"Issue\s+\d+", line, re.IGNORECASE) and re.search(
        r"(?:AESO|GASO|DHO|DHE|DHI)\s*\d{3,4}", line, re.IGNORECASE
    ):
        return True
    return False


def _is_bullet_line(line: str) -> bool:
    """Detect simple bullet or numbered list lines."""
    return bool(re.match(r"^(\d+\.|[-*•])\s+", line))


def _join_paragraph_lines(lines: Sequence[str]) -> str:
    """Merge wrapped lines into a single paragraph."""
    merged = " ".join(part.strip() for part in lines if part.strip())
    merged = re.sub(r"\s{2,}", " ", merged)
    return merged.strip()


def _split_table_block(text: str, target_chars: int) -> List[str]:
    """Split a table block only on row boundaries."""
    rows = [row.strip() for row in text.splitlines() if row.strip()]
    if not rows:
        return []
    chunks: List[str] = []
    buffer: List[str] = []
    buffer_len = 0
    for row in rows:
        row_len = len(row)
        if buffer and buffer_len + row_len + 1 > target_chars:
            chunks.append("\n".join(buffer).strip())
            buffer = []
            buffer_len = 0
        buffer.append(row)
        buffer_len += row_len + 1
    if buffer:
        chunks.append("\n".join(buffer).strip())
    return chunks


def _split_text_block(text: str, target_chars: int, overlap: int) -> List[str]:
    """Split a non-table text block while keeping boundaries natural."""
    paragraphs = [p.strip() for p in re.split(r"\n{2,}", text) if p.strip()]
    if not paragraphs:
        return []

    parts: List[str] = []
    for para in paragraphs:
        if len(para) <= target_chars:
            parts.append(para)
        else:
            parts.extend(_split_long_paragraph(para, target_chars, overlap))
    return parts


def _page_numbers(block: TextBlock) -> List[int]:
    """Return page numbers covered by a block."""
    if block.page_start is None and block.page_end is None:
        return []
    if block.page_start is None:
        return [block.page_end]  # type: ignore[list-item]
    if block.page_end is None:
        return [block.page_start]
    if block.page_end < block.page_start:
        return [block.page_start]
    return list(range(block.page_start, block.page_end + 1))


def _split_long_paragraph(text: str, target_chars: int, overlap: int) -> List[str]:
    """Split a long paragraph with sentence-aware windows, then overlap fallback."""
    cleaned = " ".join(text.split()).strip()
    if not cleaned:
        return []
    if len(cleaned) <= target_chars:
        return [cleaned]

    sentences = _split_sentences(cleaned)
    if len(sentences) > 1:
        sentence_chunks: List[str] = []
        current = ""
        for sent in sentences:
            candidate = f"{current} {sent}".strip() if current else sent
            if len(candidate) <= target_chars:
                current = candidate
                continue

            if current:
                sentence_chunks.append(current)
                if overlap > 0:
                    tail = current.split()[-max(1, overlap // 8) :]
                    current = f"{' '.join(tail)} {sent}".strip()
                else:
                    current = sent
            else:
                sentence_chunks.append(sent)
                current = ""

        if current:
            sentence_chunks.append(current)

        if all(len(chunk) <= target_chars * 1.2 for chunk in sentence_chunks):
            logger.debug(
                "Split long paragraph into %s sentence-based segments.",
                len(sentence_chunks),
            )
            return sentence_chunks

    step = max(1, target_chars - overlap)
    segments: List[str] = []
    for start in range(0, len(cleaned), step):
        end = min(len(cleaned), start + target_chars)
        segment = cleaned[start:end].strip()
        if segment:
            segments.append(segment)
        if end == len(cleaned):
            break
    logger.debug(
        "Split long paragraph into %s overlapping segments.",
        len(segments),
    )
    return segments


def _split_sentences(text: str) -> List[str]:
    """Lightweight sentence splitter to keep chunk breaks natural."""
    parts = re.split(r"(?<=[.!?])\s+(?=[A-Z0-9])", text)
    return [p.strip() for p in parts if p and p.strip()]


def _extract_order_code(text: str | None) -> str | None:
    """Extract an order code like DHO/AESO/GASO from a heading if present."""
    if not text:
        return None
    match = re.search(r"\b(RA|AESO|GASO|DHO|DHE|DHI)\s*([0-9]{2,4})\b", text, re.IGNORECASE)
    if not match:
        return None
    prefix, digits = match.groups()
    return f"{prefix.upper()} {digits}"


def _extract_page_marker(line: str) -> int | None:
    """Detect page marker lines emitted during ingestion joins."""
    match = re.match(r"^\[\[\[PAGE_BREAK_(\d+)\]\]\]$", line.strip())
    if not match:
        return None
    return int(match.group(1))


def clean_title(title: str) -> str:
    """Return a display-friendly title by stripping boilerplate and placeholders.

    - Removes phrases like "UNCONTROLLED COPY WHEN PRINTED".
    - Normalises placeholder headings like "X X X" to "General".
    - Collapses repeated whitespace and trims common separators.
    """
    t = (title or "").strip()
    if not t:
        return t
    t = re.sub(r"\bUNCONTROLLED COPY WHEN PRINTED\b", "", t, flags=re.IGNORECASE)
    if re.fullmatch(r"(?i)(x\s+){2,}x", t):
        return "General"
    if re.fullmatch(r"(?:[A-Z]\s+){2,}[A-Z]", t):
        t = "".join(re.findall(r"[A-Z]", t))
    t = re.sub(r"\s{2,}", " ", t).strip(" \t-–·:·")
    return t or title


def _assert_pdf_has_text(
    pages: Sequence[Tuple[int, str]],
    text_lengths: Sequence[int],
    *,
    blank_threshold: int = 60,
    blank_ratio_limit: float = 0.7,
) -> None:
    """Raise if a PDF appears to lack an extractable text layer (image-only scans)."""
    if not pages:
        return

    total_chars = sum(text_lengths)
    blank_pages = sum(1 for length in text_lengths if length < blank_threshold)
    blank_ratio = blank_pages / len(pages)

    if total_chars == 0:
        raise ValueError("No extractable text found in PDF; run OCR before ingesting.")

    avg_chars = total_chars / len(pages)
    # Guard against long PDFs that only yield a few lines (typical of scanned docs).
    if len(pages) >= 6 and (blank_ratio >= blank_ratio_limit or avg_chars < 120):
        raise ValueError(
            "PDF appears to lack a text layer (likely scanned). Please OCR the file and retry."
        )
